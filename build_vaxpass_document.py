from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = "VaxPass_Project_Explanation_and_Screenshot_Guide.docx"
BLUE = "0C4A6E"; TEAL = "0F9B9D"; PALE = "E7F8F7"; GRID = "D9E4E9"; TEXT = "183146"; MUTED = "607789"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)
def border(cell, color=GRID):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for e in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), color); borders.append(el)
def set_cell(cell, text, bold=False, color=None):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); r = p.add_run(text); r.bold = bold; r.font.size = Pt(9); r.font.name = "Aptos"; r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos"); r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)
def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
    for i,h in enumerate(headers): shade(table.rows[0].cells[i], BLUE); set_cell(table.rows[0].cells[i], h, True, "FFFFFF")
    for n,row in enumerate(rows):
        cells = table.add_row().cells
        for i,val in enumerate(row):
            if n % 2 == 1: shade(cells[i], "F4F9FB")
            set_cell(cells[i], str(val))
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table
def para(doc, text, bold_lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7); p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE); p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    for r in p.runs: r.font.name="Aptos"; r._element.rPr.rFonts.set(qn("w:ascii"),"Aptos"); r._element.rPr.rFonts.set(qn("w:hAnsi"),"Aptos"); r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(TEXT)
    return p
def bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(3); r=p.add_run(item); r.font.name="Aptos"; r.font.size=Pt(10.5); r.font.color.rgb=RGBColor.from_string(TEXT)
def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before=Pt(14 if level==1 else 9); p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.font.name="Aptos Display"; r.font.color.rgb=RGBColor(0,0,0); r.font.size=Pt(16 if level==1 else 12); r.bold=True
    return p
def screenshot_box(doc, label):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; c=t.cell(0,0); shade(c,"F4F9FB"); border(c, "AFCED4"); c.height=Inches(.55)
    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(label); r.italic=True; r.font.color.rgb=RGBColor.from_string(MUTED); r.font.size=Pt(9)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles; styles["Normal"].font.name="Aptos"; styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"),"Aptos"); styles["Normal"].font.size=Pt(10.5)

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(72); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("VaxPass"); r.font.name="Aptos Display"; r.font.size=Pt(34); r.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Blockchain Based Cold Chain Vaccine Passport and Temperature Aware Expiry Management DApp"); r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(TEAL); r.bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(34); r=p.add_run("Project Explanation and Output Screenshot Guide"); r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(TEXT)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(95); r=p.add_run("College Mini Project Report Companion"); r.italic=True; r.font.color.rgb=RGBColor.from_string(MUTED)
doc.add_page_break()

heading(doc,"1 Introduction")
para(doc,"VaxPass is an Ethereum based application for tracking vaccine batches from manufacture to administration. It solves a cold chain problem: many vaccines must remain within a specified temperature range during transport and storage. If a batch becomes too warm or too cold, its usable life should reduce, and a severe breach should prevent its use.")
para(doc,"The project combines a smart contract, a browser interface, a MetaMask wallet, and a simulated sensor. The important idea is that the blockchain is the shared source of truth. Manufacturers, distributors, clinics, regulators, and temperature loggers see the same history, while only approved wallets can perform sensitive actions.")
heading(doc,"2 What the Project Demonstrates")
bullets(doc,["Role based access control: only approved manufacturers, distributors, clinics, temperature loggers, and regulators may perform their permitted actions.","Traceability: every batch receives a numerical blockchain ID, displayed as VX001, VX002, and so on in the frontend.","Cold chain monitoring: temperature logs are stored against a batch and out of range readings consume an excursion budget.","Temperature aware expiry: each excursion shortens the effective expiry date by a deterministic formula.","Safety controls: recalled, expired, and spoiled batches cannot proceed through normal workflow actions.","Patient privacy: the application hashes a patient ID in the browser. The original patient identity is never written to the blockchain.","Public verification: anyone can verify a batch without connecting a wallet."])
heading(doc,"3 Technology Stack")
add_table(doc,["Technology","Purpose in VaxPass","Simple explanation"],[
["Solidity 0.8.20","Smart contract language","Defines the permanent rules and data stored on Ethereum."],
["Ethereum and Hardhat","Blockchain and local development network","Hardhat provides a local Ethereum chain for safe development, testing, and demonstration."],
["OpenZeppelin AccessControl","Security library","Provides trusted role management instead of writing role logic from scratch."],
["JavaScript and ethers v6","Frontend blockchain connection","Reads contract data and asks MetaMask to sign transactions."],
["HTML and CSS","User interface","Builds the responsive pages, forms, cards, timeline, and visual status badges."],
["MetaMask","Wallet","Connects a user account to the DApp and confirms on chain transactions."],
["qrcode.js and Canvas","QR and charts","Generates a public verification QR code and a simple temperature line chart."],
["Node sensor script","Sensor simulation","Sends normal, spike, or breach temperature readings using the logger account."]
],[1.35,2.25,3.1])
heading(doc,"4 System Architecture")
para(doc,"The architecture has four layers. The user interacts with the browser frontend. The frontend reads public data through a JSON RPC provider and sends write transactions through MetaMask. The Solidity contract validates every action and stores the trusted batch state. The sensor simulator represents the external temperature device by sending signed readings from a logger wallet.")
add_table(doc,["Layer","Main responsibility","Example"],[
["Frontend","Shows data and collects user input","Clinic enters a patient ID; JavaScript hashes it before sending."],
["MetaMask","Signs user approved transactions","Manufacturer confirms the transaction to register a batch."],
["VaxPass contract","Applies the business rules","Rejects a dose if the batch is spoiled or patient hash was already used."],
["Ethereum network","Stores the shared audit trail","All permitted users can inspect the same batch history."],
["Sensor simulator","Produces logger transactions","Breach mode creates long high temperature excursions."]
],[1.25,2.5,2.95])
doc.add_page_break()

heading(doc,"5 Smart Contract Workflow")
heading(doc,"5.1 Roles",2)
add_table(doc,["Role","Main permission"],[
["Regulator","Approves or revokes participants and recalls unsafe batches."],
["Manufacturer","Registers batches and dispatches them to an approved distributor."],
["Distributor","Accepts custody, chooses the destination clinic, and confirms handover."],
["Clinic","Confirms receipt and administers doses from batches in its custody."],
["Logger","Records cold chain temperature readings while a batch is in transit or delivered."]
],[1.5,5.2])
para(doc,"A user cannot simply claim a role. The wallet first requests a role, then the regulator approves it. This demonstrates access control because the smart contract checks the wallet role before each sensitive function.")
heading(doc,"5.2 Batch lifecycle",2)
para(doc,"A new batch starts as MANUFACTURED. The manufacturer dispatches it, making it IN TRANSIT. The distributor accepts custody. The distributor then assigns a clinic, and both the distributor and clinic call confirmReceipt. Only after both confirmations does the contract change the batch to DELIVERED and transfer custody to the clinic. When a clinic administers doses, the status becomes IN USE and finally DEPLETED when no doses remain.")
add_table(doc,["Status","Meaning","Normal next step"],[
["MANUFACTURED","Created by an approved manufacturer","Dispatch to distributor"],
["IN TRANSIT","Moving or waiting for handover confirmation","Accept custody, log temperature, deliver to clinic"],
["DELIVERED","Dual confirmation completed; clinic owns it","Log temperature or administer doses"],
["IN USE","At least one dose has been given","Administer remaining doses"],
["DEPLETED","No doses remain","Terminal state"],
["SPOILED, RECALLED, EXPIRED","Unsafe or no longer usable","Terminal state; actions are blocked"]
],[1.4,2.65,2.65])
heading(doc,"5.3 Why dual confirmation matters",2)
para(doc,"The delivery is not treated as complete when only the distributor says it delivered the batch. The clinic must also confirm receipt. This two party confirmation reduces disputes: the blockchain records who confirmed, when they confirmed, and the exact state change.")
heading(doc,"6 Temperature Aware Expiry Management")
para(doc,"Each batch has a permitted temperature range, for example 2.0 to 8.0 degrees Celsius, and an excursion budget measured in minutes. A reading inside the range is logged but changes nothing. A reading outside the range adds its duration to excursionMinsUsed.")
para(doc,"The contract uses a penalty factor of 10. For every one minute outside the temperature range, the effective expiry is reduced by 10 minutes. For example, a 4 minute excursion reduces effective expiry by 40 minutes. The code prevents timestamp underflow. When the total excursion time becomes greater than the allowed budget, the status automatically becomes SPOILED.")
para(doc,"This is a transparent teaching model, not a real medical calculation. Real vaccine stability rules depend on the vaccine product and require validated scientific data.")
heading(doc,"7 Privacy and Security")
bullets(doc,["The contract stores a bytes32 patient hash, not a patient name, Aadhaar number, phone number, or medical record.","The same patient hash cannot receive two doses from the same batch, preventing accidental duplicate administration for that batch.","OpenZeppelin AccessControl protects role restricted functions.","The regulator can recall a batch. Anyone may mark an active batch expired after its effective expiry time.","Every important state change emits an event and also writes a readable history item for the Passport timeline."])
doc.add_page_break()

heading(doc,"8 Frontend Pages and What to Explain During a Demo")
add_table(doc,["Page","What it does","What to say in the viva or presentation"],[
["Home","Shows recent batches and quick links","This is the dashboard and entry point for each participant."],
["Register Role","Requests an operational role","The organisation name stays in browser storage; only the requested blockchain role is sent on chain."],
["Regulator","Approves, revokes, and recalls","Regulatory control is enforced by the contract, not only by the UI."],
["Manufacturer","Creates and dispatches batches","A new on chain batch receives its traceable VX ID."],
["Distributor","Accepts custody and delivers to clinic","This represents transport accountability and dual confirmation."],
["Clinic","Receives batches and administers doses","The patient ID is hashed locally before the transaction."],
["Passport","Shows fields, logs, chart, timeline, and QR","This is the complete digital passport for one batch."],
["Verify","Public VALID, WARNING, or INVALID result","Anyone can verify a batch without a wallet or account."],
["Analytics","Shows totals and wastage","This gives a simple operational overview from on chain statistics."]
],[1.25,2.25,3.2])
heading(doc,"9 Recommended Live Demonstration")
bullets(doc,["Open Verify without MetaMask and enter VX001. This proves that public verification is read only and open to anyone.","Open the Passport page. Point out printed expiry, effective expiry, temperature range, custody, excursion bar, chart, timeline, and QR code.","Run normal sensor readings. Explain that readings within 2 to 8 degrees Celsius do not change expiry.","Run spike readings. Show that an out of range temperature increases excursion minutes and reduces effective expiry.","Run breach readings. Show the automatic SPOILED status and an INVALID verification result.","Use VX002 for the safe workflow: manufacturer dispatches, distributor accepts, distributor selects clinic, both parties confirm receipt, and the clinic receives custody.","Administer a dose from VX002. Explain the patient hash. Try the same patient ID again to show duplicate prevention.","Recall a batch as regulator, then show that Verify returns INVALID and normal actions are blocked.","Open Analytics last to show the final totals and wastage."])
heading(doc,"10 Testing Evidence")
para(doc,"The automated Hardhat test suite contains 13 passing tests. It covers role requests, approval and revocation, regulator transfer, manufacturer only registration, custody flow, normal and excursion temperature readings, automatic spoilage, delivery and dose blocking, patient hash duplicate protection, depletion, recall, expiry, invalid transitions, public verification codes, histories, logs, and aggregate statistics.")
para(doc,"Useful commands for the report are: npx hardhat compile, npx hardhat test, npm run node, npm run deploy:local, and npm run seed:local. The README contains the complete setup and Sepolia deployment instructions.")
doc.add_page_break()

heading(doc,"11 Output Screenshots Required for the Project Report")
para(doc,"Use the following screenshots as the output evidence section of the report. Keep the browser URL bar out of the crop if possible, do not show private keys or seed phrases, and give every image a short caption such as Figure 4 Cold Chain Excursion Warning.")
shots=[
("1","Home dashboard","Connected wallet, recent VX batches, and blue teal medical interface.","Shows the DApp is running and reading live blockchain data."),
("2","Public verification valid","Verify VX001 or VX002 before a breach and show the green VALID banner.","Shows public read only verification without a wallet."),
("3","Role based wallet view","Top right connected wallet and its detected Manufacturer, Distributor, Clinic, or Regulator role.","Proves the application recognises role based access."),
("4","Batch registration","Manufacturer registration form completed or a newly created batch card.","Shows creation of a traceable vaccine batch."),
("5","Custody transfer","Distributor page showing an incoming request or accepted custody.","Shows supply chain handover control."),
("6","Dual confirmation delivery","Distributor handover confirmation and/or Clinic receipt confirmation state.","Shows two party delivery confirmation."),
("7","Normal temperature logs","Passport chart/table with 2 to 8 degree readings and zero excursion use.","Shows safe cold chain operation."),
("8","Temperature spike warning","Passport with increased excursion minutes and lower effective expiry.","Shows temperature aware expiry calculation."),
("9","Spoiled and invalid batch","SPOILED badge plus red INVALID verify banner after breach mode.","Shows automatic safety enforcement."),
("10","Dose administration","Clinic page immediately before or after administering a dose; doses remaining changes.","Shows controlled clinical use."),
("11","QR code and passport timeline","Batch Passport QR code beside timeline and temperature chart.","Shows end to end traceability and easy public access."),
("12","Regulator recall","Regulator panel recall form or recalled batch with red status.","Shows regulator intervention."),
("13","Analytics","Analytics cards and bar chart after the demo.","Shows on chain operational statistics and wastage."),
("14","Automated test output","Terminal showing 13 passing tests.","Shows contract logic was tested, not only visually demonstrated.")]
add_table(doc,["No","Screenshot","Capture","Report purpose"],shots,[.4,1.35,2.75,1.35])
heading(doc,"12 Suggested Figure Captions")
bullets(doc,["Figure 1 VaxPass Home Dashboard with Recent Vaccine Batches","Figure 2 Public Blockchain Verification of Vaccine Batch VX001","Figure 3 Manufacturer Registration of a Cold Chain Vaccine Batch","Figure 4 Distributor and Clinic Dual Confirmation of Delivery","Figure 5 Temperature Excursion Reducing Effective Expiry","Figure 6 Automatic Spoilage After Excursion Budget Breach","Figure 7 Privacy Preserving Dose Administration Using Patient Hash","Figure 8 VaxPass Analytics and Wastage Statistics","Figure 9 Hardhat Automated Smart Contract Test Results"])
heading(doc,"13 Limitations and Future Scope")
para(doc,"The sensor readings in this mini project are simulated by a Node.js script. A production deployment would require trusted hardware sensors, signed device data, secure key management, reliable oracle infrastructure, product specific stability rules, privacy review, and formal security auditing. The project demonstrates the blockchain design and workflow, not a certified medical deployment.")
para(doc,"Possible future extensions include IPFS document uploads, QR labels printed on physical vaccine boxes, real IoT gateway integration, multiple regulator organisations, alert notifications, location tracking, and dashboards for hospitals or governments.")

# footer
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run("VaxPass Project Explanation and Screenshot Guide"); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

doc.core_properties.title="VaxPass Project Explanation and Screenshot Guide"
doc.core_properties.subject="College mini project documentation"
doc.core_properties.author="VaxPass Project Team"
doc.save(OUT)
print(OUT)
